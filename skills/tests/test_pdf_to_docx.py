import json
import subprocess
import sys
from pathlib import Path

import docx
import pytest

from conftest import SKILL_DIR
from pdf_to_docx import PdfPasswordError, convert, has_text_layer

SCRIPT = str(SKILL_DIR / "scripts" / "pdf_to_docx.py")


def run_cli(*args):
    return subprocess.run([sys.executable, SCRIPT, *args], capture_output=True, text=True)


def test_has_text_layer_true_for_text_pdf(pdf_chu):
    assert has_text_layer(pdf_chu) is True


def test_has_text_layer_false_for_scan(pdf_scan):
    assert has_text_layer(pdf_scan) is False


def test_fallback_keeps_vietnamese(pdf_chu, tmp_path):
    r = convert(pdf_chu, str(tmp_path / "o.docx"), force_fallback=True)
    assert r["method"] == "fallback" and r["pages"] == 2 and r["tables"] >= 1
    text = "\n".join(p.text for p in docx.Document(r["out_path"]).paragraphs)
    assert "Khu công nghiệp Long Thành" in text


def test_fallback_table_cells(pdf_chu, tmp_path):
    d = docx.Document(convert(pdf_chu, str(tmp_path / "o.docx"), force_fallback=True)["out_path"])
    cells = {c.text.strip() for t in d.tables for row in t.rows for c in row.cells}
    assert {"Diện tích", "12.500 m²"} <= cells


def test_pdf2docx_branch(pdf_chu, tmp_path):
    pytest.importorskip("pdf2docx")
    assert convert(pdf_chu, str(tmp_path / "o.docx"))["method"] == "pdf2docx"


def test_auto_fallback_when_pdf2docx_missing(pdf_chu, tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "pdf2docx", None)
    assert convert(pdf_chu, str(tmp_path / "o.docx"))["method"] == "fallback"


def test_password_raises(pdf_mat_khau, tmp_path):
    with pytest.raises(PdfPasswordError):
        convert(pdf_mat_khau, str(tmp_path / "o.docx"))


def test_cli_default_output_name(pdf_chu):
    p = run_cli(pdf_chu)
    assert p.returncode == 0, p.stderr
    out = Path(pdf_chu).with_name("bao-cao_word.docx")
    assert out.exists()
    assert json.loads(p.stdout)["is_scan"] is False


def test_cli_check_scan(pdf_scan):
    p = run_cli(pdf_scan, "--check")
    assert p.returncode == 0, p.stderr
    assert json.loads(p.stdout) == {"is_scan": True, "pages": 1}


def test_cli_password_exit_2(pdf_mat_khau):
    p = run_cli(pdf_mat_khau)
    assert p.returncode == 2
    assert "mật khẩu" in p.stderr and "Traceback" not in p.stderr
    assert not Path(pdf_mat_khau).with_name("hop-dong-khoa_word.docx").exists()


def test_cli_missing_file_exit_1(tmp_path):
    p = run_cli(str(tmp_path / "khong-co.pdf"))
    assert p.returncode == 1 and "Traceback" not in p.stderr


def test_cli_not_pdf_exit_1(tep_khong_phai_pdf):
    p = run_cli(tep_khong_phai_pdf)
    assert p.returncode == 1 and "không phải PDF" in p.stderr
