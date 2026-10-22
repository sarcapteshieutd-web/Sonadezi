"""Chuyển PDF dạng chữ sang Word (.docx) cho skill Xử lý PDF SZL.

Cách dùng:
    python pdf_to_docx.py <vao.pdf> [<ra.docx>] [--fallback] [--check]

- Mặc định ghi <tên-gốc>_word.docx cạnh tệp vào.
- Ưu tiên thư viện pdf2docx (giữ bố cục); không có thì tự dùng pdfplumber + python-docx.
- --check: chỉ kiểm tra tệp là bản scan hay có lớp chữ.
- In kết quả dạng JSON. Mã thoát: 0 thành công, 1 lỗi tệp, 2 PDF có mật khẩu.
"""
import argparse
import json
import sys
from pathlib import Path


class PdfPasswordError(Exception):
    """PDF được bảo vệ bằng mật khẩu."""


def _check_pdf(pdf_path):
    path = Path(pdf_path)
    if not path.is_file():
        raise FileNotFoundError(f"Không tìm thấy tệp: {pdf_path}")
    with open(path, "rb") as f:
        if f.read(5) != b"%PDF-":
            raise ValueError(f"Tệp không phải PDF: {pdf_path}")
    _reader(str(path))


def _reader(pdf_path):
    """Mở PDF; PDF chỉ hạn chế quyền (mật khẩu mở rỗng) vẫn đọc được."""
    from pypdf import PdfReader
    reader = PdfReader(pdf_path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise PdfPasswordError(f"PDF có mật khẩu: {pdf_path}")
    return reader


def _page_count(pdf_path):
    return len(_reader(pdf_path).pages)


def scan_pages(pdf_path: str, min_chars_per_page: int = 20) -> list:
    """Số thứ tự (từ 1) các trang gần như không có lớp chữ, tức trang scan."""
    import pdfplumber
    with pdfplumber.open(pdf_path) as pdf:
        return [i for i, p in enumerate(pdf.pages, 1)
                if len((p.extract_text() or "").strip()) < min_chars_per_page]


def has_text_layer(pdf_path: str, min_chars_per_page: int = 20) -> bool:
    import pdfplumber
    with pdfplumber.open(pdf_path) as pdf:
        if not pdf.pages:
            return False
        total = sum(len((p.extract_text() or "").strip()) for p in pdf.pages)
        return total / len(pdf.pages) >= min_chars_per_page


def _convert_pdf2docx(pdf_path, out_path):
    from pdf2docx import Converter
    cv = Converter(pdf_path)
    try:
        cv.convert(out_path)
    finally:
        cv.close()
    import docx
    return len(docx.Document(out_path).tables)


def _outside(bboxes):
    def keep(obj):
        x0, top, x1, bottom = obj["x0"], obj["top"], obj["x1"], obj["bottom"]
        return not any(x0 >= b[0] - 1 and x1 <= b[2] + 1 and top >= b[1] - 1 and bottom <= b[3] + 1
                       for b in bboxes)
    return keep


def _convert_fallback(pdf_path, out_path):
    import docx
    import pdfplumber
    document = docx.Document()
    tables = 0
    with pdfplumber.open(pdf_path) as pdf:
        for i, page in enumerate(pdf.pages):
            if i:
                document.add_page_break()
            found = page.find_tables()
            text = page.filter(_outside([t.bbox for t in found])).extract_text() or ""
            for line in text.splitlines():
                if line.strip():
                    document.add_paragraph(line.strip())
            for t in found:
                rows = [r for r in t.extract() if any(c for c in r)]
                if not rows:
                    continue
                table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        table.cell(r, c).text = (value or "").strip()
                tables += 1
    document.save(out_path)
    return tables


def convert(pdf_path: str, out_path: str, force_fallback: bool = False) -> dict:
    _check_pdf(pdf_path)
    method = "fallback"
    tables = None
    if not force_fallback:
        try:
            tables = _convert_pdf2docx(pdf_path, out_path)
            method = "pdf2docx"
        except ImportError:
            tables = None
        except Exception:  # pdf2docx lỗi với bố cục lạ: bỏ tệp dở dang, chuyển sang cách dự phòng
            Path(out_path).unlink(missing_ok=True)
            tables = None
    if tables is None:
        tables = _convert_fallback(pdf_path, out_path)
    return {"method": method, "pages": _page_count(pdf_path), "tables": tables, "out_path": out_path}


def main(argv=None):
    parser = argparse.ArgumentParser(description="Chuyển PDF sang Word (.docx)")
    parser.add_argument("pdf")
    parser.add_argument("out", nargs="?")
    parser.add_argument("--fallback", action="store_true", help="bỏ qua pdf2docx")
    parser.add_argument("--check", action="store_true", help="chỉ kiểm tra bản scan")
    args = parser.parse_args(argv)
    try:
        _check_pdf(args.pdf)
        scans = scan_pages(args.pdf)
        pages = _page_count(args.pdf)
        if args.check:
            result = {"is_scan": len(scans) == pages, "pages": pages, "scan_pages": scans}
        else:
            out = args.out or str(Path(args.pdf).with_name(Path(args.pdf).stem + "_word.docx"))
            if Path(out).suffix.lower() != ".docx" or Path(out).resolve() == Path(args.pdf).resolve():
                raise ValueError(f"Tệp đầu ra phải có đuôi .docx và khác tệp PDF gốc: {out}")
            result = convert(args.pdf, out, force_fallback=args.fallback)
            result.update({"is_scan": len(scans) == pages, "scan_pages": scans})
    except PdfPasswordError:
        print("Lỗi: PDF được bảo vệ bằng mật khẩu. Đề nghị gỡ mật khẩu rồi gửi lại tệp.", file=sys.stderr)
        return 2
    except FileNotFoundError as e:
        print(f"Lỗi: {e}", file=sys.stderr)
        return 1
    except ValueError as e:
        print(f"Lỗi: {e}", file=sys.stderr)
        return 1
    except Exception as e:  # tệp hỏng hoặc lỗi thư viện: báo gọn, không traceback
        print(f"Lỗi: không xử lý được tệp ({e.__class__.__name__}: {e})", file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
